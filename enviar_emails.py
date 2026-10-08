import os
import smtplib
import time
import mimetypes
from datetime import datetime
from email.message import EmailMessage

from openpyxl import load_workbook

from dotenv import load_dotenv
load_dotenv()

REMETENTE = "antonio.carosendo@gmail.com"
SENHA = os.environ["GMAIL_APP_PASSWORD"]
PLANILHA_FUNCIONARIOS = "funcionarios.xlsx"
ANEXO = "painel_vendas.xlsx"
ASSUNTO = "Dashboard de vendas da papelaria - 2024"
PAUSA_ENTRE_ENVIOS = 2

MENSAGEM = """Olá, {nome}!

Segue o dashboard de vendas da papelaria referente a 2024, para análise
de insights do setor de {setor}.

Atenciosamente,
Antonio Carlos
"""

def ler_contatos(caminho):
    planilha = load_workbook(caminho)
    aba = planilha.active
    contatos = []
    for nome, email, setor in aba.iter_rows(min_row=2, values_only=True):
        if email:
            contatos.append({"nome": nome, "email": email, "setor": setor})
    return contatos


def montar_email(contato):
    msg = EmailMessage()
    msg["From"] = REMETENTE
    msg["To"] = contato["email"]
    msg["Subject"] = ASSUNTO
    msg.set_content(MENSAGEM.format(nome=contato["nome"], setor=contato["setor"]))

    if ANEXO:
        tipo, _ = mimetypes.guess_type(ANEXO)
        tipo = tipo or "application/octet-stream"
        maintype, subtype = tipo.split("/")
        with open(ANEXO, "rb") as arquivo:
            msg.add_attachment(
                arquivo.read(),
                maintype=maintype,
                subtype=subtype,
                filename=os.path.basename(ANEXO),
            )
    return msg


def main():
    funcionarios = ler_contatos(PLANILHA_FUNCIONARIOS)
    print(f"{len(funcionarios)} contatos encontrados.")

    enviados, falhas = 0, []

    planilha = load_workbook(PLANILHA_FUNCIONARIOS)

    if "Registros" in planilha.sheetnames:
        aba_registro = planilha["Registros"]
    else:
        aba_registro = planilha.create_sheet("Registros")
        aba_registro.append(["Nome", "Email", "Data", "Status"])

    #Conexão única para todos os envios
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as servidor:
        servidor.login(REMETENTE, SENHA)

        for funcionario in funcionarios:
            try:
                servidor.send_message(montar_email(funcionario))
                enviados += 1
                print(f"Enviado para {funcionario['nome']} <{funcionario['email']}>")
                status = 'Enviado com sucesso'
            except Exception as erro:
                falhas.append(funcionario["email"])
                print(f"✘ Falha com {funcionario['email']}: {erro}")
                status = 'Falha no envio'

            data = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
            aba_registro.append([funcionario['nome'], funcionario['email'], data, status])
            time.sleep(PAUSA_ENTRE_ENVIOS)

    planilha.save(PLANILHA_FUNCIONARIOS)

    print(f"\nConcluído: {enviados} enviados, {len(falhas)} falhas.")
    if falhas:
        print("Falharam:", ", ".join(falhas))


if __name__ == "__main__":
    main()